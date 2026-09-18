"""
build_master_documentation.py
Generates the complete Ringers Backend Master Documentation (Phases 1 to 8):
1. Updates RINGERS_BACKEND_PHASE_1_TO_7_MASTER_DOCUMENTATION.md
2. Compiles Ringers_Backend_Master_Documentation_Phase1_to_7.docx
"""

import os
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOCS_DIR = os.path.abspath(os.path.dirname(__file__))
MD_PATH = os.path.join(DOCS_DIR, "RINGERS_BACKEND_PHASE_1_TO_7_MASTER_DOCUMENTATION.md")
DOCX_PATH = os.path.join(DOCS_DIR, "Ringers_Backend_Master_Documentation_Phase1_to_7.docx")
DOCX_P8_PATH = os.path.join(DOCS_DIR, "Ringers_Backend_Master_Documentation_Phase1_to_8.docx")

def build_markdown():
    # Read existing markdown file
    with open(MD_PATH, "r", encoding="utf-8") as f:
        content = f.read()

    # Update title and metadata
    content = content.replace(
        "# Ringers Platform — Master Technical Documentation\n### Complete Architectural & Engineering Reference: Phases 1 to 7\n**Version:** 2.0.0 | **Date:** September 2026 | **Repository:** [https://github.com/gayatrii-ii/ringers-react](https://github.com/gayatrii-ii/ringers-react)  \n**Current Branch:** `main` (Merged & Verified: Commit `af65b5c`)",
        "# Ringers Platform — Master Technical Documentation\n### Complete Architectural & Engineering Reference: Phases 1 to 8 (Production-Ready)\n**Version:** 2.5.0 | **Date:** September 2026 | **Repository:** [https://github.com/gayatrii-ii/ringers-react](https://github.com/gayatrii-ii/ringers-react)  \n**Current Branch:** `main` (Merged & Verified: Commit `386e021` via PR #9)"
    )

    # Update Table of Contents
    old_toc_p7 = "   - [Phase 7: Multi-Language Notifications (i18n EN/HI/MR), Customer Reviews & Analytics](#phase-7-multi-language-notifications-i18n-enhimr-customer-reviews--analytics)"
    new_toc_p8 = old_toc_p7 + "\n   - [Phase 8: Operational Flows, Customer Onboarding & Delivery Execution](#phase-8-operational-flows-customer-onboarding--delivery-execution)\n     - [8.1 Customer Dual-Flow Onboarding (Flow A & Flow B)](#81-customer-dual-flow-onboarding)\n     - [8.2 Customer-Specific Products & Custom Pricing Engine](#82-customer-specific-products--custom-pricing-engine)\n     - [8.3 Vendor-Only Delivery Dispatch & Fleet Handover Lifecycle](#83-vendor-only-delivery-dispatch--fleet-handover-lifecycle)\n     - [8.4 Connected Delivery Boy Job Application Pipeline](#84-connected-delivery-boy-job-application-pipeline)\n     - [8.5 Vendor Referral Program & Rewards](#85-vendor-referral-program--rewards)\n     - [8.6 Support & Issue Ticket System](#86-support--issue-ticket-system)\n     - [8.7 Migration V17: Schema Extensions & Indexes](#87-migration-v17-schema-extensions--indexes)"
    content = content.replace(old_toc_p7, new_toc_p8)

    # Update TOC counts
    content = content.replace(
        "5. [Complete Master API Reference (55+ Endpoints)](#5-complete-master-api-reference)\n6. [Automated Test Verification Report (251 / 251 Tests)](#6-automated-test-verification-report)",
        "5. [Complete Master API Reference (77+ Endpoints)](#5-complete-master-api-reference)\n6. [Automated Test Verification Report (329 / 329 Tests)](#6-automated-test-verification-report)"
    )

    # Update Tech Stack table
    content = content.replace(
        "| **Database Migrations** | Flyway-compatible SQL | 16 versioned migrations (`V1__` to `V16__`) |",
        "| **Database Migrations** | Flyway-compatible SQL | 17 versioned migrations (`V1__` to `V17__`) |"
    )
    content = content.replace(
        "| **Real-time / Push** | Socket.IO & FCM | Real-time order dispatch, driver GPS tracking, FCM device token management |",
        "| **Real-time / Push** | Socket.IO (v4.8) & FCM | Real-time order lifecycle events, location beacons, FCM push alerts |"
    )
    content = content.replace(
        "| **Testing** | `tsx` runner | 7 custom verification suites, 251 assertions covering 100% of business logic |",
        "| **Testing** | `tsx` runner | 8 custom verification suites, 329 assertions covering 100% of business logic |"
    )

    # Update Directory Structure
    old_dir_scripts = """│   ├── scripts/                           ← Automated phase test suites
│   │   ├── test-auth-flow.ts              ← Phase 2 verification (10 tests)
│   │   ├── test-vendor-catalog-flow.ts    ← Phase 3A verification (16 tests)
│   │   ├── test-admin-governance-flow.ts  ← Phase 3B verification (23 tests)
│   │   ├── test-customer-delivery-flow.ts ← Phase 4 verification (22 tests)
│   │   ├── test-order-flow.ts             ← Phase 5 verification (23 tests)
│   │   ├── test-payment-flow.ts           ← Phase 6 verification (46 tests)
│   │   └── test-phase7-flow.ts            ← Phase 7 verification (111 tests)"""

    new_dir_scripts = """│   ├── scripts/                           ← Automated phase test suites
│   │   ├── test-auth-flow.ts              ← Phase 2 verification (10 tests)
│   │   ├── test-vendor-catalog-flow.ts    ← Phase 3A verification (16 tests)
│   │   ├── test-admin-governance-flow.ts  ← Phase 3B verification (23 tests)
│   │   ├── test-customer-delivery-flow.ts ← Phase 4 verification (22 tests)
│   │   ├── test-order-flow.ts             ← Phase 5 verification (23 tests)
│   │   ├── test-payment-flow.ts           ← Phase 6 verification (46 tests)
│   │   ├── test-phase7-flow.ts            ← Phase 7 verification (111 tests)
│   │   └── test-phase8-flow.ts            ← Phase 8 verification (78 tests)"""
    content = content.replace(old_dir_scripts, new_dir_scripts)

    content = content.replace(
        "│   │   │   ├── reviews/                   ← Customer ratings, review masking & statistics\n│   │   │   └── vendors/                   ← Vendor storefronts, hours, addresses, staff, UPI",
        "│   │   │   ├── reviews/                   ← Customer ratings, review masking & statistics\n│   │   │   ├── support/                   ← Customer & vendor support ticketing engine\n│   │   │   └── vendors/                   ← Vendor storefronts, customer pricing, referrals, delivery onboarding"
    )

    content = content.replace(
        "│   │   │   ├── review.routes.ts           ← /api/v1/reviews/*\n│   │   │   └── vendor.routes.ts           ← /api/v1/vendors/*",
        "│   │   │   ├── review.routes.ts           ← /api/v1/reviews/*\n│   │   │   ├── support.routes.ts          ← /api/v1/support/*\n│   │   │   └── vendor.routes.ts           ← /api/v1/vendors/*"
    )

    content = content.replace(
        "│   └── db/migration/                      ← 16 Flyway SQL migrations (V1 to V16)",
        "│   └── db/migration/                      ← 17 Flyway SQL migrations (V1 to V17)"
    )

    # Phase 8 Detailed Section to insert right after Phase 7
    phase8_section = """
---

### Phase 8: Operational Flows, Customer Onboarding & Delivery Execution

**Status:** Completed, Audited & Merged  
**Goal:** Implement all remaining operational workflows required for end-to-end production operations connecting 10+ vendors, 20+ delivery partners, and 10,000+ customers: customer-specific custom pricing, dual-flow customer onboarding, vendor-only delivery dispatch, dual delivery completion (OTP & customer confirmation), delivery partner job pipeline, vendor referral program, support ticketing, and Socket.IO real-time tracking.

#### 8.1 Customer Dual-Flow Onboarding
The platform supports two distinct customer onboarding channels tailored for local neighborhood commerce:
1. **Flow A (Vendor-Initiated Onboarding):**
   - **Step 1:** Merchant registers regular customer details (first name, last name, phone, email, photo URL, delivery address, notes) via `POST /api/v1/vendors/customers/register`.
   - **Step 2:** System creates `customer.registration_requests` record with `initiated_by = 'VENDOR'` and status `PENDING`.
   - **Step 3:** System generates and dispatches a 4-digit SMS OTP to customer's mobile number.
   - **Step 4:** Merchant verifies OTP via `POST /api/v1/vendors/customers/:requestId/verify-otp`. Request transitions to `OTP_VERIFIED`.
   - **Step 5:** Merchant configures customer-specific product availability & custom pricing, sets initial password, and activates account via `POST /api/v1/vendors/customers/:requestId/activate`.
   - **Step 6 (Atomic Activation):** With transaction lock (`FOR UPDATE`), system creates `identity.users` with role `CUSTOMER`, creates `customer.customer_profiles`, default delivery address in `customer.addresses`, initializes wallet with ₹0.00 in `payment.wallets`, associates custom product pricing in `vendor.customer_products`, and marks request `ACTIVATED`.

2. **Flow B (Customer-Direct Public Registration Request):**
   - **Step 1:** Prospective customer selects a vendor and submits registration request via public endpoint `POST /api/v1/customers/registration-requests`.
   - **Step 2:** System checks phone is not already registered, creates request record (`initiated_by = 'CUSTOMER'`, `status = 'PENDING'`), and notifies vendor store owner.
   - **Step 3:** Merchant reviews pending requests on store dashboard via `GET /api/v1/vendors/customers/registration-requests`.
   - **Step 4:** Merchant activates the account and provides credentials via `POST /api/v1/vendors/customers/:requestId/activate` or rejects with reason via `PATCH /api/v1/vendors/customers/registration-requests/:requestId/reject`.

#### 8.2 Customer-Specific Products & Custom Pricing Engine
- **Per-Customer Pricing (`vendor.customer_products`):** Vendors can enable/disable individual catalog products for specific customers and configure customized unit prices (e.g. ₹42.00/L for Customer A vs ₹46.00/L catalog rate).
- **Cart & Checkout Calculation Integration:** Modified `OrderService.calculateOrder` to fetch customer product configuration. When calculating line items:
  - If custom price is configured: `effectiveUnitPrice = custProductConfig.customPrice`.
  - Discount is computed automatically if regular price > custom price.
- **Vendor Catalog Restriction Mode (`restrict_customer_catalog`):**
  - Configurable boolean toggle on `vendor.vendors` (`PATCH /api/v1/vendors/:vendorId/catalog-restriction`).
  - When enabled, customers can **only** browse and order products that have been explicitly enabled for their account by the merchant. Any attempt to order unassigned products returns `403 PRODUCT_NOT_PERMITTED`.

#### 8.3 Vendor-Only Delivery Dispatch & Fleet Handover Lifecycle
- **Strict Vendor-Only Assignment Authority:**
  - `POST /api/v1/orders/:id/assign-delivery` enforces that **only** the merchant owning the order can assign delivery riders. Super Admin and third parties are blocked with `403 VENDOR_ONLY_ACTION`.
  - Verifies rider exists, has `DELIVERY_BOY` role, and is currently `ONLINE`.
  - Atomically transitions order `delivery_status = 'ASSIGNED'`, creates `delivery.delivery_assignments` in `ASSIGNED` status, and emits real-time `assignment:created` event.
- **Rider Assignment Acceptance & Decline:**
  - `POST /api/v1/delivery/assignments/:id/accept`: Rider accepts. Assignment & order `delivery_status` transition to `ACCEPTED`. Rider status switches to `BUSY`. Real-time `assignment:accepted` emitted.
  - `POST /api/v1/delivery/assignments/:id/reject`: Rider declines with mandatory reason. Assignment marked `REJECTED`, order `delivery_status` resets to `UNASSIGNED`, rider remains `ONLINE`. Merchant is immediately alerted via WebSocket and push notification to reassign.
- **Order Pickup Handover:**
  - `POST /api/v1/delivery/assignments/:id/pickup`: Rider marks package collected at store. Assignment transitions to `PICKED_UP`, order transitions to `OUT_FOR_DELIVERY`. Automatically triggers 4-digit `DELIVERY_CONFIRMATION` OTP to customer phone and emits `order:out_for_delivery`.
- **Dual Delivery Completion Channels:**
  - **Path A (Rider submits customer OTP):** `POST /api/v1/delivery/orders/:id/complete-delivery`. Verifies 4-digit OTP. Atomically marks assignment & order `DELIVERED`, resets rider duty to `ONLINE`, and triggers automated 90% vendor payout.
  - **Path B (Customer direct app confirmation):** `POST /api/v1/orders/:id/confirm-delivery`. Customer directly clicks "Confirm Delivery Received" in mobile app. Atomically completes order, resets rider to `ONLINE`, and triggers vendor payout.
- **Delivery Failure Reporting:**
  - `POST /api/v1/delivery/assignments/:id/fail`: Rider reports delivery failure with standardized codes: `CUSTOMER_UNAVAILABLE`, `WRONG_ADDRESS`, `CUSTOMER_REFUSED`, `CANNOT_CONTACT_CUSTOMER`, `OTHER`.
  - Assignment set to `CANCELLED`, order set to `CANCELLED` with `delivery_status = 'FAILED'`, rider reset to `ONLINE`, and merchant notified.
- **Live GPS Tracking & Location Beacons:**
  - `POST /api/v1/delivery/location`: Rider sends GPS coordinates (latitude, longitude, accuracy). Saved to append-only `delivery.delivery_locations` table and emitted via WebSocket `location:update` to tracking customer.
  - `GET /api/v1/delivery/track/:id`: Unified tracking endpoint returning order status, delivery partner info, vehicle type, and latest GPS coordinate.

#### 8.4 Connected Delivery Boy Job Application Pipeline
- **Public Submission:** Prospective riders apply via `POST /api/v1/public/delivery-job-request` (name, phone, vehicle type: `BIKE`, `SCOOTER`, `CYCLE`, `ELECTRIC_VEHICLE`, license number).
- **Admin Fleet Connection:** Super Admin reviews applicants via `GET /api/v1/admin/delivery-boy-requests` and connects rider to a merchant fleet (`status = 'CONNECTED'`).
- **Vendor Storefront Activation:** Connected merchants view requests via `GET /api/v1/vendors/delivery-boys/job-requests`, provision login credentials, and activate account via `POST /api/v1/vendors/delivery-boys/activate`.

#### 8.5 Vendor Referral Program & Rewards
- **Referral Code Engine:** Vendors obtain a unique referral code (`REF-XXXX-XXXX`) and shareable invite link (`GET /api/v1/vendors/referral`).
- **Invitations:** Merchants invite prospective merchants via phone or email (`POST /api/v1/vendors/referrals/invite`).
- **Reward Lifecycle:** `PENDING` → `APPROVED` → `PAID` / `CANCELLED`.
- **Governance:** Super Admin monitors platform referrals (`GET /api/v1/admin/referrals`) and approves reward amounts (`PATCH /api/v1/admin/referrals/:id/reward`).

#### 8.6 Support & Issue Ticket System
- **Ticket Creation:** Authenticated customers, vendors, and riders submit support tickets via `POST /api/v1/support/tickets` (categories: `ORDER_ISSUE`, `PAYMENT_ISSUE`, `DELIVERY_ISSUE`, `GENERAL`; priorities: `LOW`, `MEDIUM`, `HIGH`, `URGENT`).
- **Collision-Proof Ticket Numbers:** Automatically generated as `RNG-TCK-YYYYMMDD-XXXX`.
- **User Dashboard:** Users view ticket history via `GET /api/v1/support/tickets/my`.
- **Admin Resolution Workflow:** Super Admin triages tickets via `GET /api/v1/support/admin/tickets` and resolves with recorded response via `PATCH /api/v1/support/admin/tickets/:id/resolve`.

#### 8.7 Migration V17: Schema Extensions & Indexes
- **Migration:** `V17__create_phase8_operational_tables.sql`
- **Tables Created / Extended:**
  1. `support.tickets`: Customer and vendor support tickets with status and priority constraints.
  2. `vendor.customer_products`: Per-customer product enablement and custom price overrides with unique constraint `uq_vendor_customer_product (vendor_id, customer_user_id, product_id)`.
  3. `customer.registration_requests`: Dual-flow registration tracking with `initiated_by IN ('VENDOR', 'CUSTOMER')`.
  4. `vendor.referrals`: Referral program tracking with referee contact, status, and reward tracking.
  5. `vendor.vendors`: Extended with `restrict_customer_catalog BOOLEAN DEFAULT FALSE` and `referral_code VARCHAR(50) UNIQUE`.
  6. `delivery.delivery_boy_job_requests`: Extended with `status IN ('PENDING', 'CONNECTED', 'ACTIVATED', 'REJECTED')` and `activated_user_id`.
  7. `delivery.delivery_assignments`: Extended with `rejection_reason`, `failure_reason`, `failure_notes`, `rejected_at`, `failed_at`.
  8. `order_management.orders`: Extended `delivery_status` check constraint to include `ACCEPTED`.
"""

    content = content.replace("--- \n\n## 5. Complete Master API Reference", phase8_section + "\n---\n\n## 5. Complete Master API Reference")
    if phase8_section not in content:
        content = content.replace("--- \r\n\r\n## 5. Complete Master API Reference", phase8_section + "\n---\n\n## 5. Complete Master API Reference")
    if phase8_section not in content:
        content = content.replace("---\n\n## 5. Complete Master API Reference", phase8_section + "\n---\n\n## 5. Complete Master API Reference")

    # Add Phase 8 API Endpoints to Section 5
    old_vendor_apis = """| `PATCH`| `/vendors/:id/status` | `SUPER_ADMIN` | Approve, activate, or block vendor account |
| `PUT` | `/vendors/:id` | `SUPER_ADMIN` | Administrative override of vendor profile |"""

    new_vendor_apis = """| `PATCH`| `/vendors/:id/status` | `SUPER_ADMIN` | Approve, activate, or block vendor account |
| `PUT` | `/vendors/:id` | `SUPER_ADMIN` | Administrative override of vendor profile |
| `POST` | `/vendors/customers/register` | `VENDOR` | Flow A: Vendor initiates customer registration & triggers OTP |
| `POST` | `/vendors/customers/:requestId/verify-otp` | `VENDOR` | Flow A: Vendor verifies customer SMS OTP |
| `POST` | `/vendors/customers/:requestId/activate` | `VENDOR` | Flow A/B: Vendor activates customer account with password & pricing |
| `GET` | `/vendors/customers/registration-requests` | `VENDOR` | Vendor lists incoming customer registration requests |
| `PATCH`| `/vendors/customers/registration-requests/:requestId/reject` | `VENDOR` | Vendor rejects customer registration request with reason |
| `GET` | `/vendors/:vendorId/customers/:customerId/products` | `VENDOR` | View customer-specific product availability & custom pricing |
| `PUT` | `/vendors/:vendorId/customers/:customerId/products` | `VENDOR` | Batch configure customer products & custom unit pricing |
| `PATCH`| `/vendors/:vendorId/catalog-restriction` | `VENDOR` | Toggle customer catalog restriction setting ON/OFF |
| `GET` | `/vendors/delivery-boys/job-requests` | `VENDOR` | View delivery boy applications connected to vendor fleet |
| `POST` | `/vendors/delivery-boys/activate` | `VENDOR` | Activate connected delivery boy account with login credentials |
| `GET` | `/vendors/referral` | `VENDOR` | Get vendor referral dashboard profile, metrics, and invite code |
| `POST` | `/vendors/referrals/invite` | `VENDOR` | Invite prospective merchant via phone or email |
| `GET` | `/vendors/referrals` | `VENDOR` | List all referrals dispatched by vendor |"""
    content = content.replace(old_vendor_apis, new_vendor_apis)

    # Customer registration request endpoint
    old_cust_apis = "| `GET` | `/customers/profile/me` | `CUSTOMER` | View customer profile (name, email, phone, DOB, gender) |"
    new_cust_apis = "| `POST` | `/customers/registration-requests` | Public | Flow B: Customer submits direct registration request to vendor |\n" + old_cust_apis
    content = content.replace(old_cust_apis, new_cust_apis)

    # Delivery Handover endpoints
    old_del_apis = """| `GET` | `/delivery/stats` | `DELIVERY_BOY` | View delivery count, active assignment, and metrics |
| `POST` | `/delivery/location` | `DELIVERY_BOY` | Ingest live GPS coordinates (latitude, longitude, heading) |"""

    new_del_apis = """| `GET` | `/delivery/stats` | `DELIVERY_BOY` | View delivery count, active assignment, and metrics |
| `POST` | `/delivery/assignments/:id/accept` | `DELIVERY_BOY` | Rider accepts delivery assignment (duty switches to BUSY) |
| `POST` | `/delivery/assignments/:id/reject` | `DELIVERY_BOY` | Rider rejects assignment with reason (order returned to vendor) |
| `POST` | `/delivery/assignments/:id/pickup` | `DELIVERY_BOY` | Rider marks picked up (triggers delivery confirmation OTP to customer) |
| `POST` | `/delivery/orders/:id/complete-delivery` | `DELIVERY_BOY` | Path A: Rider enters customer OTP to complete delivery & trigger payout |
| `POST` | `/delivery/assignments/:id/fail` | `DELIVERY_BOY` | Rider reports delivery failure with standardized reason codes |
| `POST` | `/delivery/location` | `DELIVERY_BOY` | Ingest live GPS coordinates (latitude, longitude, heading) |
| `GET` | `/delivery/track/:id` | Authenticated | Real-time order tracking with driver info and latest GPS beacon |"""
    content = content.replace(old_del_apis, new_del_apis)

    # Order Delivery Handover endpoints
    old_ord_apis = "| `PATCH`| `/orders/:id/cancel` | Authenticated | Customer or Vendor order cancellation with reason |"
    new_ord_apis = """| `PATCH`| `/orders/:id/cancel` | Authenticated | Customer or Vendor order cancellation with reason |
| `POST` | `/orders/:id/assign-delivery` | `VENDOR` | Vendor assigns delivery rider to order (Vendor-only authority) |
| `POST` | `/orders/:id/confirm-delivery` | `CUSTOMER` | Path B: Customer directly confirms delivery received in app |"""
    content = content.replace(old_ord_apis, new_ord_apis)

    # Admin referral governance
    old_adm_apis = "| `PATCH`| `/admin/delivery-boy-requests/:id/assign` | `SUPER_ADMIN` | Approve rider and link to vendor's delivery fleet |"
    new_adm_apis = """| `PATCH`| `/admin/delivery-boy-requests/:id/assign` | `SUPER_ADMIN` | Approve rider and link to vendor's delivery fleet |
| `GET` | `/admin/referrals` | `SUPER_ADMIN` | List all merchant referrals across platform |
| `PATCH`| `/admin/referrals/:id/reward` | `SUPER_ADMIN` | Approve, pay, or cancel referral reward amount |"""
    content = content.replace(old_adm_apis, new_adm_apis)

    # Support Ticket Module
    support_api_section = """
### 12. Support & Issue Tickets Module (`/api/v1/support`)
| Method | Endpoint | Access / Role | Description |
|---|---|---|---|
| `POST` | `/support/tickets` | Authenticated | Submit new support ticket with category and priority |
| `GET` | `/support/tickets/my` | Authenticated | View authenticated user's submitted support tickets |
| `GET` | `/support/tickets/:id` | Authenticated | View single ticket details and resolution audit history |
| `GET` | `/support/admin/tickets` | `SUPER_ADMIN` / `ADMIN` | List all platform tickets with priority sorting |
| `PATCH`| `/support/admin/tickets/:id/resolve` | `SUPER_ADMIN` / `ADMIN` | Resolve or close support ticket with recorded admin response |
"""
    if "### 12. Support & Issue Tickets Module" not in content:
        content = content.replace("--- \n\n## 6. Automated Test Verification Report", support_api_section + "\n---\n\n## 6. Automated Test Verification Report")
        if "### 12. Support & Issue Tickets Module" not in content:
            content = content.replace("---\n\n## 6. Automated Test Verification Report", support_api_section + "\n---\n\n## 6. Automated Test Verification Report")

    # Update Test Table in Section 6
    old_test_table = """| **Phase 7** | Notifications (i18n EN/HI/MR), Reviews & Analytics | `test-phase7-flow.ts` | **111 / 111** | ✅ PASS |
| **TOTAL** | **Full Platform Regression Suite** | `npm run test:all` | **251 / 251** | **✅ 100% PASS** |"""

    new_test_table = """| **Phase 7** | Notifications (i18n EN/HI/MR), Reviews & Analytics | `test-phase7-flow.ts` | **111 / 111** | ✅ PASS |
| **Phase 8** | Customer Onboarding, Custom Pricing, Handover, Support | `test-phase8-flow.ts` | **78 / 78** | ✅ PASS |
| **TOTAL** | **Full Platform Regression Suite (8 Phases)** | `npm run test:all` | **329 / 329** | **✅ 100% PASS** |"""
    content = content.replace(old_test_table, new_test_table)

    # Update Git History in Section 7
    old_git_history = """| `phase-7-notifications-analytics-and-reviews` | PR #7 | Team | Phase 7: i18n notifications, Customer reviews, Analytics dashboards |
| **`main`** | — | — | **Stable Production Branch (Latest Commit: `af65b5c`)** |"""

    new_git_history = """| `phase-7-notifications-analytics-and-reviews` | PR #7 | Team | Phase 7: i18n notifications, Customer reviews, Analytics dashboards |
| `phase-8-missing-backend-delivery-and-auth-flows` | PR #9 | Karan | Phase 8: Operational flows, dual onboarding, pricing, delivery handover, referrals, support |
| **`main`** | — | — | **Stable Production Branch (Latest Commit: `386e021`)** |"""
    content = content.replace(old_git_history, new_git_history)

    # Update Checklist
    old_checklist = """- [x] **Phase 7: Notifications & Reviews:** 3-Language i18n (EN/HI/MR), FCM device tokens, 1-5★ verified reviews, Analytics.
- [x] **Automated Tests Passing:** **251 / 251 (100%)** via `npm run test:all`.
- [x] **Git Status:** Merged into `main` (Commit `af65b5c`), 0 conflicts, 0 errors."""

    new_checklist = """- [x] **Phase 7: Notifications & Reviews:** 3-Language i18n (EN/HI/MR), FCM device tokens, 1-5★ verified reviews, Analytics.
- [x] **Phase 8: Operational Flows:** Dual customer onboarding, custom product pricing, delivery handover (OTP & confirmation), rider job pipeline, vendor referrals, support tickets.
- [x] **Automated Tests Passing:** **329 / 329 (100%)** via `npm run test:all`.
- [x] **Git Status:** Merged into `main` (Commit `386e021`), 0 conflicts, 0 errors."""
    content = content.replace(old_checklist, new_checklist)

    with open(MD_PATH, "w", encoding="utf-8") as f:
        f.write(content)
    print("Updated markdown saved successfully.")
    return content

# Helper functions for styling python-docx
def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'''<w:tcMar {nsdecls("w")}>
        <w:top w:w="{top}" w:type="dxa"/>
        <w:bottom w:w="{bottom}" w:type="dxa"/>
        <w:left w:w="{left}" w:type="dxa"/>
        <w:right w:w="{right}" w:type="dxa"/>
    </w:tcMar>''')
    tcPr.append(tcMar)

def set_table_borders(table):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(f'''<w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
            <w:insideV w:val="none"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
        </w:tblBorders>''')
        tblPr[0].append(borders)

def build_docx(markdown_content):
    doc = Document()

    # Configure Margins: 0.75 in
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(30, 41, 59) # Slate-800
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(4)

    # Title Banner Block
    title_p = doc.add_paragraph()
    title_run = title_p.add_run("RINGERS QUICK-COMMERCE PLATFORM")
    title_run.font.name = 'Arial'
    title_run.font.size = Pt(24)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(15, 23, 42) # Slate-900

    sub_p = doc.add_paragraph()
    sub_run = sub_p.add_run("Master Technical Architecture & Engineering Documentation — Phases 1 to 8")
    sub_run.font.name = 'Arial'
    sub_run.font.size = Pt(14)
    sub_run.font.bold = True
    sub_run.font.color.rgb = RGBColor(37, 99, 235) # Blue-600

    meta_p = doc.add_paragraph()
    meta_run = meta_p.add_run("Version 2.5.0 (Production Verified)  |  September 2026  |  Repository: github.com/gayatrii-ii/ringers-react\nBranch: main (Merged Commit: 386e021)  |  Engineering Team: Karan Khot, Gayatri, Karansinh, Pradnya, Gauri")
    meta_run.font.size = Pt(9.5)
    meta_run.font.color.rgb = RGBColor(100, 116, 139) # Slate-500
    meta_p.paragraph_format.space_after = Pt(16)

    # Divider
    p_div = doc.add_paragraph()
    p_div_run = p_div.add_run("━" * 58)
    p_div_run.font.color.rgb = RGBColor(226, 232, 240)

    # Parse Markdown lines into Document elements
    lines = markdown_content.split("\n")
    i = 0
    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    def flush_table(t_lines):
        if not t_lines:
            return
        rows = [row.strip().strip("|").split("|") for row in t_lines if not re.match(r'^\s*\|?\s*[-:]+[-| :]*$', row)]
        if not rows:
            return
        num_cols = max(len(r) for r in rows)
        tbl = doc.add_table(rows=len(rows), cols=num_cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl)

        for r_idx, row in enumerate(rows):
            is_header = (r_idx == 0)
            tr = tbl.rows[r_idx]
            for c_idx, val in enumerate(row):
                if c_idx < num_cols:
                    cell = tr.cells[c_idx]
                    clean_text = val.strip().replace('**', '').replace('`', '').replace('\\', '')
                    cell.text = clean_text
                    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.05
                    run = p.runs[0] if p.runs else p.add_run(clean_text)
                    run.font.name = 'Calibri'
                    if is_header:
                        set_cell_background(cell, "1E3A8A") # Navy Blue
                        run.font.bold = True
                        run.font.size = Pt(9.5)
                        run.font.color.rgb = RGBColor(255, 255, 255)
                    else:
                        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
                        set_cell_background(cell, bg_color)
                        run.font.size = Pt(9)
                        run.font.color.rgb = RGBColor(30, 41, 59)
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    def flush_code(c_lines):
        if not c_lines:
            return
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.rows[0].cells[0]
        set_cell_background(cell, "F1F5F9") # Slate-100
        set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run("\n".join(c_lines))
        run.font.name = 'Consolas'
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(15, 23, 42)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Skip original title block
    while i < len(lines) and not lines[i].startswith("## 📑 Table of Contents") and not lines[i].startswith("## 1. Executive Summary"):
        i += 1

    while i < len(lines):
        line = lines[i]

        # Handle Code Blocks
        if line.strip().startswith("```"):
            if in_code_block:
                flush_code(code_lines)
                code_lines = []
                in_code_block = False
            else:
                if in_table:
                    flush_table(table_lines)
                    table_lines = []
                    in_table = False
                in_code_block = True
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Handle Tables
        if line.strip().startswith("|") and line.strip().endswith("|"):
            in_table = True
            table_lines.append(line)
            i += 1
            continue
        elif in_table:
            flush_table(table_lines)
            table_lines = []
            in_table = False

        stripped = line.strip()

        # Handle Headings
        if stripped.startswith("# "):
            h = doc.add_heading(level=1)
            run = h.add_run(stripped[2:])
            run.font.name = 'Arial'
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(4)

        elif stripped.startswith("## "):
            h = doc.add_heading(level=2)
            run = h.add_run(stripped[3:])
            run.font.name = 'Arial'
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 58, 138) # Dark Blue
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)

        elif stripped.startswith("### "):
            h = doc.add_heading(level=3)
            run = h.add_run(stripped[4:])
            run.font.name = 'Arial'
            run.font.size = Pt(11.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(37, 99, 235) # Accent Blue
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)

        elif stripped.startswith("#### "):
            h = doc.add_heading(level=4)
            run = h.add_run(stripped[5:])
            run.font.name = 'Calibri'
            run.font.size = Pt(10.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(71, 85, 105) # Slate-600
            h.paragraph_format.space_before = Pt(6)
            h.paragraph_format.space_after = Pt(2)

        elif stripped.startswith("> "):
            # Callout Quote / Alert Box
            callout = stripped[2:].replace('[!IMPORTANT]', 'IMPORTANT:').replace('[!NOTE]', 'NOTE:')
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.rows[0].cells[0]
            set_cell_background(cell, "EFF6FF") # Light Blue
            set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(callout)
            run.font.italic = True
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(30, 64, 175)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)

        elif stripped.startswith("- [x]") or stripped.startswith("- [ ]"):
            # Checklist
            p = doc.add_paragraph(style='List Bullet')
            is_checked = stripped.startswith("- [x]")
            prefix = "☑ " if is_checked else "☐ "
            text = stripped[5:].strip()
            run_check = p.add_run(prefix)
            run_check.font.bold = True
            run_check.font.color.rgb = RGBColor(16, 185, 129) if is_checked else RGBColor(156, 163, 175)
            # Add formatted runs
            add_formatted_runs(p, text)
            p.paragraph_format.space_after = Pt(2)

        elif stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            text = stripped[2:].strip()
            add_formatted_runs(p, text)
            p.paragraph_format.space_after = Pt(2)

        elif re.match(r'^\d+\.\s+', stripped):
            p = doc.add_paragraph()
            add_formatted_runs(p, stripped)
            p.paragraph_format.space_after = Pt(2)

        elif stripped == "---":
            pass

        elif stripped:
            p = doc.add_paragraph()
            add_formatted_runs(p, stripped)

        i += 1

    if in_table:
        flush_table(table_lines)

    # Save to DOCX_PATH and DOCX_P8_PATH
    doc.save(DOCX_PATH)
    doc.save(DOCX_P8_PATH)
    print(f"Saved DOCX to: {DOCX_PATH}")
    print(f"Saved DOCX to: {DOCX_P8_PATH}")

def add_formatted_runs(paragraph, text):
    parts = re.split(r'(\*\*.*?\*\*|\`.*?\`|\*.*?\*)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            r = paragraph.add_run(part[2:-2])
            r.font.bold = True
        elif part.startswith('`') and part.endswith('`'):
            r = paragraph.add_run(part[1:-1])
            r.font.name = 'Consolas'
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(190, 24, 93) # Pink/Purple code
        elif part.startswith('*') and part.endswith('*'):
            r = paragraph.add_run(part[1:-1])
            r.font.italic = True
        else:
            paragraph.add_run(part)

if __name__ == "__main__":
    md_content = build_markdown()
    build_docx(md_content)
